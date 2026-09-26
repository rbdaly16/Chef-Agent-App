"""Chef Agent API: Extract, translate, and chat about recipes."""

from __future__ import annotations

import io
import json
import os
import re
from pathlib import Path

import httpx
from dotenv import load_dotenv
from fastapi import FastAPI, Query
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from agent import Recipe, ask_chef, extract_recipes, translate_recipe

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DATA_PATH = ROOT / "data"
DOCX_PATH = DATA_PATH / "recipes.docx"

load_dotenv(ROOT / ".env")
load_dotenv(ROOT.parent / ".env")

app = FastAPI(title="Chef Agent", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def _load_recipes_from_file() -> list[Recipe]:
    """Load recipes from the docx file or return empty list."""
    if not DOCX_PATH.exists():
        return []
    try:
        from docx import Document
        doc = Document(DOCX_PATH)
        raw_text = "\n".join(p.text for p in doc.paragraphs)
        return extract_recipes(raw_text)
    except Exception as e:
        print(f"Error loading recipes: {e}")
        return []


def _load_recipes_from_google_doc() -> list[Recipe]:
    """Load recipes from a linked Google Doc if URL is set."""
    url = os.environ.get("RECIPES_GOOGLE_DOC_URL", "").strip()
    if not url:
        return []
    try:
        match = re.search(r"/d/([a-zA-Z0-9_-]+)", url)
        doc_id = match.group(1) if match else url
        export = f"https://docs.google.com/document/d/{doc_id}/export?format=docx"
        response = httpx.get(export, timeout=30, follow_redirects=True)
        if response.status_code in (401, 403, 404):
            return []
        response.raise_for_status()
        if not response.content.startswith(b"PK"):
            return []
        from docx import Document
        doc = Document(io.BytesIO(response.content))
        raw_text = "\n".join(p.text for p in doc.paragraphs)
        return extract_recipes(raw_text)
    except Exception as e:
        print(f"Error loading from Google Doc: {e}")
        return []


_recipes_cache: list[Recipe] | None = None


def load_recipes() -> list[Recipe]:
    """Load recipes, checking Google Doc first, then local file."""
    global _recipes_cache
    if _recipes_cache is not None:
        return _recipes_cache
    _recipes_cache = _load_recipes_from_google_doc()
    if not _recipes_cache:
        _recipes_cache = _load_recipes_from_file()
    return _recipes_cache


class ExtractRequest(BaseModel):
    text: str = Field(min_length=1)


class ExtractResponse(BaseModel):
    recipes: list[Recipe]


class TranslateRequest(BaseModel):
    recipe: Recipe


class TranslateResponse(BaseModel):
    english: Recipe
    hindi: Recipe


class ChatRequest(BaseModel):
    question: str = Field(min_length=1)


class ChatResponse(BaseModel):
    answer: str


@app.get("/api/health")
def health():
    return {"ok": True, "recipes_count": len(load_recipes())}


@app.get("/api/recipes")
def list_recipes():
    """Return all loaded recipes."""
    recipes = load_recipes()
    return {"count": len(recipes), "recipes": recipes}


@app.post("/api/extract", response_model=ExtractResponse)
def extract(body: ExtractRequest):
    """Extract structured recipes from raw text."""
    recipes = extract_recipes(body.text)
    return ExtractResponse(recipes=recipes)


@app.post("/api/translate", response_model=TranslateResponse)
def translate(body: TranslateRequest):
    """Translate a recipe to Hindi."""
    hindi = translate_recipe(body.recipe)
    return TranslateResponse(english=body.recipe, hindi=hindi)


@app.post("/api/chat", response_model=ChatResponse)
def chat(body: ChatRequest):
    """Ask the chef a question about recipes."""
    recipes = load_recipes()
    if not recipes:
        return ChatResponse(answer="No recipes loaded. Upload or link a recipes document first.")
    answer = ask_chef(body.question, recipes)
    return ChatResponse(answer=answer)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=False)
